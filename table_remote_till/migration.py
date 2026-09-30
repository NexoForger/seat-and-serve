"""Mapped legacy export staging and cautious master-data import."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation

import frappe
from frappe.utils import today
from openpyxl import load_workbook


HISTORY_TYPES = ("Sale", "Return", "Payment", "Purchase", "Stock", "Journal", "Employee")
ALL_TYPES = ("Item", "Customer", "Supplier", *HISTORY_TYPES)
SOURCES = {"Omega POS", "BIM POS", "Squirrel Cloud", "Squirrel 11", "Generic"}
MASTER_FIELDS = {
	"Item": ("name", "item_group", "stock_uom"),
	"Customer": ("name", "customer_group", "territory"),
	"Supplier": ("name", "supplier_group"),
}
ALIASES = {
	"source_id": ("source_id", "external_id", "id", "record_id", "code", "number"),
	"name": ("name", "item_name", "customer_name", "supplier_name", "description"),
	"code": ("item_code", "sku", "barcode", "code"),
	"item_group": ("item_group", "category", "department"),
	"stock_uom": ("stock_uom", "uom", "unit"),
	"is_stock_item": ("is_stock_item", "stock_item", "inventory_item"),
	"customer_group": ("customer_group", "group", "type"),
	"supplier_group": ("supplier_group", "group", "type"),
	"territory": ("territory", "region", "area"),
	"date": ("posting_date", "transaction_date", "date", "created_at"),
	"branch": ("branch", "store", "location", "outlet"),
	"amount": ("grand_total", "total", "amount", "net_amount"),
	"selling_rate": ("selling_rate", "selling_price", "sale_price", "retail_price", "unit_price", "price"),
	"currency": ("currency", "currency_code", "price_currency"),
}

# BIM POS sells multiple product lines, so these are suggestions only. Every
# client reviews the actual exported headers before importing.
BIM_ALIASES = {
	"Item": {
		"source_id": ("prodnum", "product_id", "product_number", "plu"),
		"name": ("prodname", "product_name", "product_description"),
		"code": ("product_code", "prodcode", "sku", "barcode"),
		"item_group": ("category_name", "category", "department_name", "family"),
		"stock_uom": ("unit_of_measure", "uom", "unit"),
		"is_stock_item": ("stockable", "has_stock", "is_stock_item"),
		"selling_rate": ("saleprice", "selling_price", "price", "unit_price"),
		"currency": ("currency", "currency_code", "price_currency"),
	},
	"Customer": {
		"source_id": ("custnum", "customer_id", "customer_number", "client_id"),
		"name": ("custname", "customer_name", "client_name"),
		"customer_group": ("customer_group", "customer_type", "group_name"),
		"territory": ("territory", "area", "zone", "region"),
	},
	"Supplier": {
		"source_id": ("vendornum", "vendor_id", "supplier_id", "supplier_number"),
		"name": ("vendor_name", "supplier_name"),
		"supplier_group": ("vendor_group", "supplier_group", "group_name"),
	},
	"History": {
		"source_id": ("transaction_id", "invoice_number", "receipt_number", "check_number", "ticket_number"),
		"date": ("posting_date", "transaction_date", "business_date", "date"),
		"branch": ("branch_name", "branch_id", "store", "location"),
		"amount": ("grand_total", "total_amount", "amount", "total"),
	},
}


def _manager():
	if not set(frappe.get_roles()).intersection({"System Manager", "TRT Manager"}):
		frappe.throw("Manager role required", frappe.PermissionError)


@frappe.whitelist(methods=["POST"])
def new_job(source, record_type, target_price_list=None):
	"""Create an import job directly from the onboarding wizard."""
	_manager()
	if source not in SOURCES:
		frappe.throw("Choose a supported source")
	if record_type not in ALL_TYPES:
		frappe.throw("Choose a supported record type")
	if target_price_list:
		if record_type != "Item" or not frappe.db.exists("Price List", {
			"name": target_price_list, "selling": 1, "enabled": 1}):
			frappe.throw("Choose an enabled selling Price List for Item prices")
	job = frappe.get_doc({"doctype": "TRT Import Job", "source": source,
		"record_type": record_type, "target_price_list": target_price_list,
		"status": "Uploaded"}).insert(ignore_permissions=True)
	return {"job": job.name, "import_supported": record_type in MASTER_FIELDS}


@frappe.whitelist(methods=["POST"])
def attach_export(job_name, file_name):
	"""Accept only a private CSV/XLSX attached to this exact import job."""
	_manager()
	job = frappe.get_doc("TRT Import Job", job_name)
	file = frappe.get_doc("File", file_name)
	if not file.is_private or file.attached_to_doctype != job.doctype or file.attached_to_name != job.name:
		frappe.throw("Upload a private export attached to this import job", frappe.PermissionError)
	if not file.file_url or not file.file_url.lower().endswith((".csv", ".xlsx")):
		frappe.throw("Only CSV and XLSX exports are supported")
	if file.file_size and file.file_size > 20_000_000:
		frappe.throw("Export exceeds 20 MB; split by period")
	job.source_file = file.file_url
	job.status = "Uploaded"
	job.mapping_json = None
	job.preview_json = None
	job.report_json = None
	job.save(ignore_permissions=True)
	return {"job": job.name, "file_url": file.file_url}


def _rows(job):
	if not job.source_file:
		frappe.throw("Attach a private CSV or XLSX export")
	file_name = frappe.db.get_value("File", {"file_url": job.source_file,
		"attached_to_doctype": job.doctype, "attached_to_name": job.name,
		"is_private": 1}, "name")
	if not file_name:
		frappe.throw("Private export attached to this import job was not found")
	file = frappe.get_doc("File", file_name)
	if not file.is_private:
		frappe.throw("Legacy exports must be private")
	# Frappe's default get_content() guesses text encodings and can decode an
	# Arabic Windows-1256 CSV as a different single-byte codec. Read raw bytes.
	content = file.get_content(encodings=[])
	if isinstance(content, str):
		content = content.encode("utf-8")
	if len(content) > 20_000_000:
		frappe.throw("Export exceeds 20 MB; split by period")
	if job.source_file.lower().endswith(".csv"):
		try:
			decoded = content.decode("utf-8-sig")
		except UnicodeDecodeError:
			decoded = content.decode("cp1256")
		try:
			dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|")
		except csv.Error:
			dialect = csv.excel
		reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
		headings = list(reader.fieldnames or [])
		if len(headings) != len(set(headings)) or any(not heading.strip() for heading in headings):
			frappe.throw("Export needs unique, non-empty column headers")
		return headings, list(reader)
	if job.source_file.lower().endswith(".xlsx"):
		book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		sheet = book.active
		iterator = sheet.iter_rows(values_only=True)
		headings = [str(cell).strip() if cell is not None else "" for cell in next(iterator)]
		if len(headings) != len(set(headings)) or any(not heading for heading in headings):
			frappe.throw("Export needs unique, non-empty column headers")
		rows = [dict(zip(headings, ("" if cell is None else str(cell) for cell in row), strict=False))
			for row in iterator]
		book.close()
		return headings, rows
	frappe.throw("Only CSV and XLSX exports are supported")


def _mapped(row, mapping, field):
	return str(row.get(mapping.get(field, ""), "") or "").strip()


def _normalized_heading(value):
	value = unicodedata.normalize("NFKC", str(value)).casefold()
	return "_".join("".join(char if char.isalnum() else " " for char in value).split())


@frappe.whitelist()
def suggest_mapping(job_name):
	"""Suggest mappings from column names; the manager must review them."""
	_manager()
	job = frappe.get_doc("TRT Import Job", job_name)
	headings, _ = _rows(job)
	normalized = {_normalized_heading(heading): heading for heading in headings}
	suggestions = {}
	aliases_by_field = dict(ALIASES)
	if job.source == "BIM POS":
		presets = BIM_ALIASES.get(job.record_type, BIM_ALIASES["History"])
		aliases_by_field = {field: (*presets.get(field, ()), *aliases)
			for field, aliases in ALIASES.items()}
	for semantic, aliases in aliases_by_field.items():
		for alias in aliases:
			if alias in normalized:
				suggestions[semantic] = normalized[alias]
				break
	return {"mapping": suggestions, "columns": headings,
		"unmapped_columns": [heading for heading in headings if heading not in suggestions.values()]}


def _source_key(job, row, mapping):
	identity = "\x1f".join((job.source, job.record_type, _mapped(row, mapping, "branch"),
		_mapped(row, mapping, "source_id")))
	return hashlib.sha256(identity.encode()).hexdigest()


def _inspect(job, mapping):
	headings, rows = _rows(job)
	required = ("source_id",) + MASTER_FIELDS.get(job.record_type, ("date", "branch", "amount"))
	missing_columns = [field for field in required if mapping.get(field) not in headings]
	for optional in ("selling_rate", "currency") if job.record_type == "Item" else ():
		if mapping.get(optional) and mapping[optional] not in headings:
			missing_columns.append(optional)
	if missing_columns:
		frappe.throw("Map required columns: " + ", ".join(missing_columns))
	price_currency = None
	if job.record_type == "Item" and mapping.get("selling_rate"):
		if not job.target_price_list or not frappe.db.exists("Price List", {
			"name": job.target_price_list, "selling": 1, "enabled": 1}):
			frappe.throw("Choose an enabled target selling Price List before previewing item prices")
		price_currency = frappe.db.get_value("Price List", job.target_price_list, "currency")
	totals = defaultdict(lambda: {"count": 0, "amount": Decimal(0)})
	errors = []
	seen = set()
	for index, row in enumerate(rows, start=2):
		source_id = _mapped(row, mapping, "source_id")
		key = _source_key(job, row, mapping)
		if not source_id:
			errors.append({"row": index, "error": "Missing source ID"})
		if key in seen:
			errors.append({"row": index, "error": "Duplicate source ID in branch"})
		seen.add(key)
		for field in required:
			if not _mapped(row, mapping, field):
				errors.append({"row": index, "error": f"Missing {field}"})
		if job.record_type == "Item" and mapping.get("selling_rate") in headings:
			try:
				price = Decimal(_mapped(row, mapping, "selling_rate"))
				if not price.is_finite() or price < 0:
					raise InvalidOperation
			except InvalidOperation:
				errors.append({"row": index, "error": "Invalid selling rate"})
			if mapping.get("currency") and _mapped(row, mapping, "currency") != price_currency:
				errors.append({"row": index, "error": f"Price currency must be {price_currency}"})
		amount = Decimal(0)
		if job.record_type in HISTORY_TYPES:
			try:
				amount = Decimal(_mapped(row, mapping, "amount"))
			except InvalidOperation:
				errors.append({"row": index, "error": "Invalid amount"})
		date = _mapped(row, mapping, "date")
		period = date[:7] if len(date) >= 7 else "unknown"
		bucket = totals[(_mapped(row, mapping, "branch") or "unspecified", period)]
		bucket["count"] += 1
		bucket["amount"] += amount
	return rows, {
		"source": job.source, "record_type": job.record_type, "rows": len(rows),
		"target_price_list": job.target_price_list,
		"price_currency": price_currency,
		"import_supported": job.record_type in MASTER_FIELDS,
		"totals": [{"branch": branch, "period": period, "count": value["count"],
			"amount": str(value["amount"])} for (branch, period), value in sorted(totals.items())],
		"errors": errors[:100], "error_count": len(errors),
		"columns": headings,
	}


@frappe.whitelist()
def preview(job_name, mapping_json=None):
	_manager()
	job = frappe.get_doc("TRT Import Job", job_name)
	mapping = frappe.parse_json(mapping_json or job.mapping_json or "{}")
	if not isinstance(mapping, dict):
		frappe.throw("Mapping must be a JSON object")
	_, report = _inspect(job, mapping)
	job.mapping_json = json.dumps(mapping)
	job.preview_json = json.dumps(report)
	job.status = "Validated" if not report["error_count"] else "Mapped"
	job.save(ignore_permissions=True)
	return report


def _make_master(kind, row, mapping):
	name = _mapped(row, mapping, "name")
	if kind == "Item":
		return frappe.get_doc({"doctype": "Item",
			"item_code": _mapped(row, mapping, "code") or _mapped(row, mapping, "source_id"),
			"item_name": name, "item_group": _mapped(row, mapping, "item_group"),
			"stock_uom": _mapped(row, mapping, "stock_uom"),
			"is_stock_item": int(_mapped(row, mapping, "is_stock_item") or "0")})
	if kind == "Customer":
		return frappe.get_doc({"doctype": "Customer", "customer_name": name,
			"customer_group": _mapped(row, mapping, "customer_group"),
			"territory": _mapped(row, mapping, "territory")})
	return frappe.get_doc({"doctype": "Supplier", "supplier_name": name,
		"supplier_group": _mapped(row, mapping, "supplier_group")})


@frappe.whitelist(methods=["POST"])
def import_masters(job_name):
	_manager()
	job = frappe.get_doc("TRT Import Job", job_name)
	if job.record_type not in MASTER_FIELDS:
		frappe.throw("Historical transactions require a source-specific reconciliation importer")
	mapping = frappe.parse_json(job.mapping_json or "{}")
	rows, report = _inspect(job, mapping)
	if report["error_count"]:
		frappe.throw("Resolve preview errors before import")
	job.status = "Importing"
	job.save(ignore_permissions=True)
	outcome = Counter()
	issues = []
	for index, row in enumerate(rows, start=2):
		key = _source_key(job, row, mapping)
		if frappe.db.exists("TRT Legacy Record", {"source_key": key}):
			outcome["already_imported"] += 1
			continue
		frappe.db.savepoint(f"trt_import_{index}")
		try:
			doc = _make_master(job.record_type, row, mapping)
			doc.insert(ignore_permissions=True)
			if job.record_type == "Item" and mapping.get("selling_rate"):
				frappe.get_doc({"doctype": "Item Price", "item_code": doc.name,
					"price_list": job.target_price_list, "selling": 1,
					"price_list_rate": float(Decimal(_mapped(row, mapping, "selling_rate"))),
					"valid_from": today()}).insert(ignore_permissions=True)
			frappe.get_doc({"doctype": "TRT Legacy Record", "source": job.source,
				"record_type": job.record_type, "source_id": _mapped(row, mapping, "source_id"),
				"source_key": key, "job": job.name, "frappe_doctype": doc.doctype,
				"frappe_name": doc.name, "branch": _mapped(row, mapping, "branch"),
				"status": "Imported"}).insert(ignore_permissions=True)
			outcome["imported"] += 1
		except Exception as error:
			frappe.db.rollback(save_point=f"trt_import_{index}")
			outcome["failed"] += 1
			issues.append({"row": index, "source_id": _mapped(row, mapping, "source_id"),
				"error": str(error)[:500]})
	job.imported_count = outcome["imported"]
	job.error_count = outcome["failed"]
	job.report_json = json.dumps({"counts": dict(outcome), "errors": issues[:100]})
	job.status = "Complete" if not outcome["failed"] else "Failed"
	job.save(ignore_permissions=True)
	return frappe.parse_json(job.report_json)


@frappe.whitelist()
def source_inventory(source):
	_manager()
	if source not in SOURCES:
		frappe.throw("Choose a supported source")
	jobs = frappe.get_all("TRT Import Job", filters={"source": source},
		fields=["name", "record_type", "status", "imported_count", "error_count"])
	covered = {job.record_type for job in jobs}
	return {"jobs": jobs, "missing_export_types": [kind for kind in ALL_TYPES if kind not in covered],
		"historical_import_available": False}
